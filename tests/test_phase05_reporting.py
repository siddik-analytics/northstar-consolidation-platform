"""
Phase 5 - the governed reporting marts and the Excel management model.

Two things are being protected here.

**The marts must not drift from the consolidation.** Everything downstream reads them, so a
mart that disagrees with `fact_financials` is a report that disagrees with the ledger, and no
amount of care in the workbook can repair that.

**The workbook must not restate a definition.** Excel is allowed presentation mathematics --
subtotals, variance dollars, percentages, display logic -- and is not allowed to re-derive
EBITDA, the add-back policy, FX, eliminations, NCI, the covenant rules or the cash flow. The
test for that is not stylistic: the reported figures are read back out of the calculated
workbook and compared with the marts.
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

import duckdb
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.marts.config import CONTROL_RESULTS, DATA, DUCKDB_PATH, MANIFEST, TOL_MART_USD

WORKBOOK = DATA / "90_exports" / "Northstar_Consolidation_Management_Reporting.xlsx"
QA_RESULTS = DATA / "phase05_workbook_qa.csv"

pytestmark = pytest.mark.skipif(not DUCKDB_PATH.exists(),
                                reason="the warehouse has not been built")


@pytest.fixture(scope="module")
def con():
    c = duckdb.connect(str(DUCKDB_PATH), read_only=True)
    yield c
    c.close()


def rows(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# =====================================================================================
# the marts against the consolidation
# =====================================================================================
def test_every_reporting_control_passes():
    failed = [r for r in rows(CONTROL_RESULTS)
              if r["status"] == "FAIL" and r["severity"] == "BLOCKING"]
    assert failed == [], [r["control_id"] for r in failed]


def test_the_financial_mart_reconciles_to_the_fact(con):
    worst = con.execute("""
        WITH fact AS (
            SELECT f.period_key,
                   round(-coalesce(sum(f.amount_usd) FILTER (
                       WHERE a.fs_caption_l1 = 'Revenue'), 0), 2) AS revenue_usd
            FROM vw_statutory_fact f JOIN dim_account a USING (group_account)
            WHERE f.counts_in_result AND NOT a.is_statistical
              AND f.group_account NOT IN (SELECT group_account FROM ref_ic_side)
            GROUP BY 1
        ),
        mart AS (
            SELECT period_key, round(sum(mtd_usd), 2) AS revenue_usd
            FROM mart_financial_ytd
            WHERE measure_code = 'REVENUE' AND basis = 'STATUTORY' AND scenario_code = 'ACT'
            GROUP BY 1
        )
        SELECT max(abs(m.revenue_usd - f.revenue_usd)) FROM mart m JOIN fact f USING (period_key)
    """).fetchone()[0]
    assert float(worst) <= TOL_MART_USD


def test_the_statements_still_balance_and_tie(con):
    bs = con.execute("""SELECT max(abs(t)) FROM (SELECT period_key, sum(balance_usd) AS t
                        FROM mart_balance_sheet GROUP BY 1)""").fetchone()[0]
    cf = con.execute("""SELECT max(abs(opening_cash_usd + operating_cash_flow_usd
                                       + investing_cash_flow_usd + financing_cash_flow_usd
                                       + fx_effect_on_cash_usd - closing_cash_usd))
                        FROM mart_cash_flow""").fetchone()[0]
    assert float(bs) <= 0.01 and float(cf) <= 0.01


def test_no_reserved_scenario_is_offered(con):
    """
    The Downside case is architecture that has not been populated. Offering it hands a reader
    an empty report that looks like a real one.
    """
    assert con.execute(
        "SELECT count(*) FROM dim_report_scenario WHERE scenario_code = 'DS'").fetchone()[0] == 0
    assert con.execute(
        "SELECT count(*) FROM mart_financial_monthly WHERE scenario_code = 'DS'"
    ).fetchone()[0] == 0


def test_favourability_is_account_aware(con):
    """Revenue above plan is favourable; operating expense above plan is not."""
    wrong = con.execute("""
        SELECT count(*) FROM mart_variance
        WHERE (favourable_direction = 'HIGHER' AND var_ytd_usd > 0
               AND ytd_favourability <> 'FAVOURABLE')
           OR (favourable_direction = 'LOWER' AND var_ytd_usd > 0
               AND ytd_favourability <> 'UNFAVOURABLE')
    """).fetchone()[0]
    assert wrong == 0


def test_covenant_leverage_uses_a_twelve_month_window(con):
    """
    A fiscal-year EBITDA divided into a full net debt balance reports a breach that does not
    exist as soon as the year is in progress: at August 2026 it produced 7.38x against a 4.50x
    limit, and the twelve months to that date is 4.21x and compliant.
    """
    partial = con.execute(
        "SELECT count(*) FROM mart_covenants WHERE months_in_window <> 12").fetchone()[0]
    assert partial == 0
    breaks = con.execute("""
        SELECT m.fiscal_year FROM mart_covenants m JOIN rpt_ebitda_bridge b USING (fiscal_year)
        WHERE m.period_key % 100 = 12
          AND abs(m.covenant_ebitda_usd - b.covenant_ebitda_usd) > 0.01
    """).fetchall()
    assert breaks == []


def test_the_group_is_compliant_at_every_test_date(con):
    """
    The agreement tests at fiscal year ends. A quarterly presentation stamped BREACH on two
    dates the agreement never tests, which is the most expensive mistake a board pack can make.
    """
    breaches = con.execute("""
        SELECT period_key, net_leverage, max_net_leverage FROM mart_covenants
        WHERE period_key % 100 = 12 AND NOT in_compliance
    """).fetchall()
    assert breaches == []


# =====================================================================================
# the workbook
# =====================================================================================
def test_the_workbook_exists_and_is_not_trivial():
    assert WORKBOOK.exists()
    assert WORKBOOK.stat().st_size > 500_000


def test_the_workbook_has_every_required_sheet():
    from openpyxl import load_workbook
    wb = load_workbook(WORKBOOK, read_only=True)
    expected = ["00 Cover", "01 Executive Summary", "02 P&L", "03 Business Units",
                "04 Entities", "05 Balance Sheet", "06 Cash Flow", "07 Working Capital",
                "08 EBITDA Bridge", "09 Debt & Covenants", "10 Headcount", "11 CapEx",
                "12 FX", "13 Consolidation & Controls", "14 Variance Detail",
                "15 Data & Technical"]
    assert [s for s in wb.sheetnames if not s.startswith("_")] == expected


def test_every_data_sheet_is_hidden():
    """A reader opening the pack should see fifteen reports, not thirty-odd tabs of extracts."""
    from openpyxl import load_workbook
    wb = load_workbook(WORKBOOK)
    visible = [ws.title for ws in wb.worksheets
               if ws.title.startswith("_") and ws.sheet_state == "visible"]
    assert visible == []


def test_the_workbook_carries_no_error_values():
    """
    Read from the QA run, which opened the workbook in Excel and forced a full calculation.
    openpyxl writes formulas and does not evaluate them, so a workbook full of `#REF!` looks
    perfect to the library that produced it.
    """
    assert QA_RESULTS.exists(), "run python -m src.excel.qa"
    errors = [r for r in rows(QA_RESULTS) if r["check"] == "error values"]
    assert errors == []


def test_the_workbook_layout_passes():
    blocking = [r for r in rows(QA_RESULTS) if r["severity"] == "BLOCKING"]
    assert blocking == [], [(r["sheet"], r["check"]) for r in blocking]


def test_the_workbook_reconciles_to_the_marts():
    recon = [r for r in rows(QA_RESULTS) if r["kind"] == "reconciliation"]
    assert len(recon) >= 15
    failed = [r for r in recon if "FAIL" in r["detail"]]
    assert failed == [], [r["check"] for r in failed]


def test_the_workbook_reports_actual_only_to_the_reporting_date():
    """
    The consolidation generates Actual rows for the whole of FY2026 and the group has closed
    eight months of it. Charting the rest as actual showed the company falling off a cliff
    every September.
    """
    from src.excel.data import REPORT_PERIOD
    assert REPORT_PERIOD == 202608


# =====================================================================================
# determinism
# =====================================================================================
def test_the_mart_build_is_reproducible_from_its_inputs():
    from src.marts.run import build_id
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    assert manifest["build_id"] == build_id()


def test_the_workbook_build_is_reproducible(tmp_path):
    """
    An .xlsx is a zip and a zip records the wall-clock time of each member, so two builds of
    identical data differed for no reason anyone would care about. Timestamps and document
    properties are now fixed, and the build is a pure function of the marts.

    The QA step then opens the file in Excel and saves the calculated values back, which
    rewrites the archive. That is intended and Excel does not do it byte-identically, so the
    claim is made where it can be: on the build.
    """
    import hashlib
    manifest = json.loads((DATA / "phase05_workbook_manifest.json").read_text(encoding="utf-8"))
    from src.excel.build import build
    rebuilt = build(tmp_path / "rebuild.xlsx")
    assert hashlib.sha256(rebuilt.read_bytes()).hexdigest() == manifest["build_digest"]


def test_the_workbook_names_the_marts_it_was_built_on():
    manifest = json.loads((DATA / "phase05_workbook_manifest.json").read_text(encoding="utf-8"))
    marts = json.loads(MANIFEST.read_text(encoding="utf-8"))
    assert manifest["reporting_mart_build_id"] == marts["build_id"]
    assert manifest["consolidation_build_id"] == marts["consolidation_build_id"]


def test_the_marts_are_built_on_the_frozen_consolidation():
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    consol = json.loads((DATA / "phase04_manifest.json").read_text(encoding="utf-8"))
    assert manifest["consolidation_build_id"] == consol["build_id"]
    assert manifest["source_layer_digest"] == consol["source_layer_digest"]
