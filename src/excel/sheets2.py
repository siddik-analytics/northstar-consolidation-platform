"""
Report sheets 03 to 15.

Split from `sheets.py` only for length. Every layout rule stated there applies here.
"""

from __future__ import annotations

from openpyxl.chart import Reference
from openpyxl.styles import Font

from . import style as S
from .data import ACT_VERSION, BUD_VERSION, FC_VERSION, PY_VERSION, REPORT_FY, REPORT_PERIOD
from .sheets import (column_group, marker, ACTUAL_MONTHS, FIRST_DATA_COL, MONTHS, _colour_variance, bar_chart,
                     chart_slots, col, headers, label_style, line_chart, measure_style,
                     note, put, section, std_widths, title)


# ------------------------------------------------------------------ 03 Business units
def business_units(wb, meta):
    ws = wb.create_sheet("03 Business Units")
    S.sheet_setup(ws, freeze="C8", zoom=100)
    std_widths(ws, label_width=32, data_width=12, n_data=9)

    title(ws, "Business unit performance",
          f"Year to date {meta['report_label']} · USD millions · statutory basis",
          width_cols=11)

    section(ws, 5, "Year to date", last_col=11, right_text="Actual vs budget and prior year")
    bu_hdr = 7
    headers(ws, 7, ["Actual", "Budget", "Var $", "Var %", "Prior year", "Var $",
                    "Share / margin", "Budget", "Movement"], label_text="Business unit")
    # Three analytical columns on a table of reported ones: the budget variance pair, the
    # prior-year variance, and the movement in share. Washed individually, because each sits
    # beside the figure it is derived from.
    column_group(ws, bu_hdr, FIRST_DATA_COL + 2, FIRST_DATA_COL + 3)
    column_group(ws, bu_hdr, FIRST_DATA_COL + 5, FIRST_DATA_COL + 5)
    column_group(ws, bu_hdr, FIRST_DATA_COL + 8, FIRST_DATA_COL + 8)
    r = 8
    for measure, mname in (("REVENUE", "Revenue"), ("ADJ_EBITDA", "Adjusted EBITDA")):
        put(ws, f"B{r}", mname, "ns_label_sub")
        r += 1
        first = r
        for bu, bu_name in meta["bus"]:
            put(ws, f"B{r}", bu_name, "ns_label")
            ws[f"C{r}"] = (f'=SUMIFS(bu_ytd,bu_measure,"{measure}",bu_bu,"{bu}",'
                           f'bu_ver,"{ACT_VERSION}",bu_period,{REPORT_PERIOD})')
            ws[f"D{r}"] = (f'=SUMIFS(bu_ytd,bu_measure,"{measure}",bu_bu,"{bu}",'
                           f'bu_ver,"{BUD_VERSION}",bu_period,{REPORT_PERIOD})')
            ws[f"E{r}"] = f"=C{r}-D{r}"
            ws[f"F{r}"] = f'=IF(D{r}=0,"",E{r}/ABS(D{r}))'
            ws[f"G{r}"] = (f'=SUMIFS(bu_ytd,bu_measure,"{measure}",bu_bu,"{bu}",'
                           f'bu_ver,"{PY_VERSION}",bu_period,{REPORT_PERIOD})')
            ws[f"H{r}"] = f"=C{r}-G{r}"
            # Revenue divided by revenue is 100% for every unit and says nothing. The revenue
            # block therefore shows each unit's SHARE of the group and the EBITDA block shows
            # its margin -- two different questions through one pair of columns.
            scope = "" if measure == "REVENUE" else f'bu_bu,"{bu}",'
            ws[f"I{r}"] = (f'=IFERROR(C{r}/SUMIFS(bu_ytd,bu_measure,"REVENUE",{scope}'
                           f'bu_ver,"{ACT_VERSION}",bu_period,{REPORT_PERIOD}),"")')
            ws[f"J{r}"] = (f'=IFERROR(D{r}/SUMIFS(bu_ytd,bu_measure,"REVENUE",{scope}'
                           f'bu_ver,"{BUD_VERSION}",bu_period,{REPORT_PERIOD}),"")')
            ws[f"K{r}"] = f'=IF(OR(I{r}="",J{r}=""),"",(I{r}-J{r})*100)'
            for c in "CDEGH":
                ws[f"{c}{r}"].style = "ns_m1"
            for c in "FIJ":
                ws[f"{c}{r}"].style = "ns_pct"
            ws[f"K{r}"].style = "ns_pp"
            r += 1
        put(ws, f"B{r}", f"Total business units", "ns_label_sub")
        for c in "CDEGH":
            ws[f"{c}{r}"] = f"=SUM({c}{first}:{c}{r-1})"
            ws[f"{c}{r}"].style = "ns_m1_sub"
        ws[f"F{r}"] = f'=IF(D{r}=0,"",E{r}/ABS(D{r}))'
        ws[f"F{r}"].style = "ns_pct_sub"
        ws[f"I{r}"] = (f'=IFERROR(C{r}/SUMIFS(bu_ytd,bu_measure,"REVENUE",'
                       f'bu_ver,"{ACT_VERSION}",bu_period,{REPORT_PERIOD}),"")')
        ws[f"J{r}"] = (f'=IFERROR(D{r}/SUMIFS(bu_ytd,bu_measure,"REVENUE",'
                       f'bu_ver,"{BUD_VERSION}",bu_period,{REPORT_PERIOD}),"")')
        ws[f"K{r}"] = f'=IF(OR(I{r}="",J{r}=""),"",(I{r}-J{r})*100)'
        for c in "IJ":
            ws[f"{c}{r}"].style = "ns_pct_sub"
        ws[f"K{r}"].style = "ns_pp"
        _colour_variance(ws, ["E", "H"], first, r)
        r += 2

    note(ws, r, "The total is the sum of the reportable business units. It is not the "
                "consolidated group result: consolidation adjustments — purchase price "
                "amortisation, unrealised profit, the minority attribution — are group-level "
                "entries and belong to no business unit. Sheet 13 reconciles the two.",
         last_col=11)
    r += 2

    section(ws, r, "Trend", last_col=11)
    slots, chart_w = chart_slots(ws, 11)
    # The unit the Executive Summary flags is the one highlighted here, in copper, on its
    # Actual bar. Copper means "the one management is looking at" -- a pointer, not a verdict:
    # the shortfall itself is stated in the table above, in the status colour it deserves.
    bu_names = [r[0] for r in wb["_chart"].iter_rows(min_row=2, max_row=6, min_col=8,
                                                     max_col=8, values_only=True)]
    highlight = ({bu_names.index(meta["flagged_bu"]): S.COPPER}
                 if meta.get("flagged_bu") in bu_names else None)
    bar_chart(ws, f"{slots[0]}{r + 2}", "Revenue by business unit — YTD actual against budget",
              Reference(wb["_chart"], min_col=8, min_row=2, max_row=6),
              [(Reference(wb["_chart"], min_col=12, min_row=2, max_row=6), "Actual YTD",
                S.ACTUAL),
               (Reference(wb["_chart"], min_col=13, min_row=2, max_row=6), "Budget YTD",
                S.BUDGET)], width=chart_w, height=7.4, page_break=True,
              point_colours=highlight)
    line_chart(ws, f"{slots[1]}{r + 2}", "Group gross margin — FY2026 by month",
               Reference(wb["_chart"], min_col=2, min_row=2, max_row=13),
               [(Reference(wb["_chart"], min_col=6, min_row=2, max_row=13), "Gross margin",
                 S.ACTUAL, None)], width=chart_w, height=7.4, number_format=S.PCT1)
    ws.print_area = f"A1:K{r + 18}"
    return ws


# ------------------------------------------------------------------ 04 Entities
def entities(wb, meta):
    ws = wb.create_sheet("04 Entities")
    S.sheet_setup(ws, freeze="C8", zoom=100)
    S.widths(ws, {"A": 2.2, "B": 44, "C": 7, "D": 7, "E": 10, "F": 11, "G": 12, "H": 11,
                  "I": 11, "J": 9, "K": 10, "L": 8, "M": 13})

    title(ws, "Legal entity performance",
          f"Year to date {meta['report_label']} · USD millions · statutory basis",
          width_cols=13)

    section(ws, 5, "By legal entity", last_col=13, right_text="Actual, year to date")
    ent_hdr = 7
    headers(ws, 7, ["BU", "Ccy", "ERP", "Revenue", "Gross profit", "EBITDA", "Net income",
                    "Margin", "Cash", "FTE", "Working capital"], label_text="Entity")
    # Every column on this table is an amount except Margin, which is a ratio derived from
    # two of them. The wash marks the one column a reader should not add down.
    column_group(ws, ent_hdr, FIRST_DATA_COL + 7, FIRST_DATA_COL + 7)
    r = 8
    first = r
    # The entities that belong to the unit the Executive Summary flags carry the same copper
    # pointer in the gutter that the unit's bar carries on the Business Units chart: one
    # subject, followed from the group page down to the legal entities that make it up.
    flagged_code = next((c for c, n in meta["bus"] if n == meta.get("flagged_bu")), None)
    for code, name, bu, ccy, erp in meta["entities"]:
        if bu == flagged_code:
            marker(ws, f"A{r}", "copper")
        put(ws, f"B{r}", f"{code}   {name}", "ns_label")
        put(ws, f"C{r}", bu, "ns_text_c")
        put(ws, f"D{r}", ccy, "ns_text_c")
        put(ws, f"E{r}", erp, "ns_text_c")
        for c, m in (("F", "REVENUE"), ("G", "GROSS_PROFIT"), ("H", "EBITDA"),
                     ("I", "NET_INCOME")):
            ws[f"{c}{r}"] = (f'=SUMIFS(ent_ytd,ent_measure,"{m}",ent_entity,"{code}",'
                             f'ent_period,{REPORT_PERIOD})')
            ws[f"{c}{r}"].style = "ns_m1"
        ws[f"J{r}"] = f'=IFERROR(H{r}/F{r},"")'
        ws[f"J{r}"].style = "ns_pct"
        ws[f"K{r}"] = (f'=SUMIFS(ent_cash,ent_measure,"REVENUE",ent_entity,"{code}",'
                       f'ent_period,{REPORT_PERIOD})')
        ws[f"K{r}"].style = "ns_m1"
        ws[f"L{r}"] = (f'=SUMIFS(ent_fte,ent_measure,"REVENUE",ent_entity,"{code}",'
                       f'ent_period,{REPORT_PERIOD})')
        ws[f"L{r}"].style = "ns_fte"
        ws[f"M{r}"] = (f'=SUMIFS(ent_wc,ent_measure,"REVENUE",ent_entity,"{code}",'
                       f'ent_period,{REPORT_PERIOD})')
        ws[f"M{r}"].style = "ns_m1"
        r += 1
    put(ws, f"B{r}", "Total, operating entities", "ns_label_sub")
    for c in "FGHIKM":
        ws[f"{c}{r}"] = f"=SUM({c}{first}:{c}{r-1})"
        ws[f"{c}{r}"].style = "ns_m1_sub"
    ws[f"L{r}"] = f"=SUM(L{first}:L{r-1})"
    ws[f"L{r}"].style = "ns_fte"
    ws[f"J{r}"] = f'=IFERROR(H{r}/F{r},"")'
    ws[f"J{r}"].style = "ns_pct_sub"
    r += 2

    note(ws, r, "Each entity's own consolidated contribution: its reported result translated "
                "to USD with the intercompany trade it conducts removed. Group-level "
                "consolidation entries are posted to the elimination entities and are "
                "excluded here, so the total is the operating group before those entries. "
                "Cash and working capital are balance sheet positions at the reporting date.",
         last_col=13)
    ws.print_area = f"A1:M{r + 1}"
    return ws


# ------------------------------------------------------------------ 05 Balance sheet
def balance_sheet(wb, meta):
    ws = wb.create_sheet("05 Balance Sheet")
    S.sheet_setup(ws, freeze="C8", zoom=100, landscape=False)
    std_widths(ws, label_width=42, data_width=13.5, n_data=6)

    title(ws, "Consolidated balance sheet",
          f"At {meta['report_label']} · USD millions · statutory basis", width_cols=8)

    section(ws, 5, "Financial position", last_col=8, right_text="Assets = liabilities + equity")
    bs_hdr = 7
    headers(ws, 7, ["Current", "Prior month", "Movement", "Prior year", "Movement",
                    "% of assets"], label_text="")
    column_group(ws, bs_hdr, FIRST_DATA_COL + 2, FIRST_DATA_COL + 2)
    column_group(ws, bs_hdr, FIRST_DATA_COL + 4, FIRST_DATA_COL + 4)
    r = 8
    open_section: list[int] = []
    subtotal_row: dict[str, int] = {}
    for caption, cls, kind in meta["bs_rows"]:
        if kind == "SECTION":
            put(ws, f"B{r}", caption, "ns_label_sub")
            open_section = [r + 1]
            r += 1
            continue
        if kind == "SUBTOTAL":
            subtotal_row[caption] = r
            put(ws, f"B{r}", caption, "ns_label_sub")
            for c in "CDF":
                ws[f"{c}{r}"] = f"=SUM({c}{open_section[0]}:{c}{r-1})"
                ws[f"{c}{r}"].style = "ns_m1_sub"
            ws[f"E{r}"] = f"=C{r}-D{r}"
            ws[f"G{r}"] = f"=C{r}-F{r}"
            for c in "EG":
                ws[f"{c}{r}"].style = "ns_m1_sub"
            ws[f"H{r}"] = f'=IFERROR(C{r}/$C${meta["bs_total_assets_row"]},"")'
            ws[f"H{r}"].style = "ns_pct_sub"
            r += 1
            continue
        if kind == "GRANDTOTAL":
            put(ws, f"B{r}", caption, "ns_label_tot")
            parts = meta["bs_grand"][caption]
            for c in "CDF":
                ws[f"{c}{r}"] = "=" + "+".join(f"{c}{subtotal_row[p]}" for p in parts)
                ws[f"{c}{r}"].style = "ns_m1_tot"
            ws[f"E{r}"] = f"=C{r}-D{r}"
            ws[f"G{r}"] = f"=C{r}-F{r}"
            for c in "EG":
                ws[f"{c}{r}"].style = "ns_m1_tot"
            ws[f"H{r}"] = f'=IFERROR(C{r}/$C${meta["bs_total_assets_row"]},"")'
            ws[f"H{r}"].style = "ns_pct_sub"
            r += 1
            continue
        put(ws, f"B{r}", caption, "ns_label")
        sign = "-" if cls in ("LIABILITY", "EQUITY") else ""
        ws[f"C{r}"] = (f'={sign}SUMIFS(bs_bal,bs_caption,"{caption}",bs_class,"{cls}",'
                       f'bs_period,{REPORT_PERIOD})')
        ws[f"D{r}"] = (f'={sign}SUMIFS(bs_prior_month,bs_caption,"{caption}",bs_class,"{cls}",'
                       f'bs_period,{REPORT_PERIOD})')
        ws[f"F{r}"] = (f'={sign}SUMIFS(bs_prior_year,bs_caption,"{caption}",bs_class,"{cls}",'
                       f'bs_period,{REPORT_PERIOD})')
        ws[f"E{r}"] = f"=C{r}-D{r}"
        ws[f"G{r}"] = f"=C{r}-F{r}"
        ws[f"H{r}"] = f'=IFERROR(C{r}/$C${meta["bs_total_assets_row"]},"")'
        for c in "CDEFG":
            ws[f"{c}{r}"].style = "ns_m1"
        ws[f"H{r}"].style = "ns_pct"
        r += 1

    r += 1
    # The accounting identity, stated as a reference: a copper key beside the check and the
    # condition it must meet in copper. The figure itself stays ink -- whether it IS nil is a
    # fact the reader takes from the number, not from a colour.
    marker(ws, f"A{r}", "copper")
    put(ws, f"B{r}", "Check — total assets less liabilities and equity", "ns_label")
    ws[f"C{r}"] = (f'=C{meta["bs_total_assets_row"]}-C{meta["bs_total_le_row"]}')
    ws[f"C{r}"].style = "ns_m2"
    put(ws, f"D{r}", "must be nil", "ns_note_copper")
    r += 2

    note(ws, r, "Goodwill and acquired intangibles arise on consolidation and exist in no "
                "entity's own books. Investments in subsidiaries and intercompany balances "
                "eliminate to nil and are shown rather than suppressed, because a caption "
                "that is nil because it eliminated is evidence and an absent caption is not.",
         last_col=8)
    ws.print_area = f"A1:H{r + 1}"
    return ws


# ------------------------------------------------------------------ 06 Cash flow
def cash_flow(wb, meta):
    ws = wb.create_sheet("06 Cash Flow")
    S.sheet_setup(ws, freeze="C8", zoom=100)
    std_widths(ws, label_width=40, data_width=10.5, n_data=9)

    title(ws, "Consolidated cash flow and liquidity",
          f"FY{REPORT_FY} to {meta['report_label']} · USD millions · "
          f"derived from balance sheet movements", width_cols=11)

    section(ws, 5, "Cash flow", last_col=11, right_text="Actual")
    headers(ws, 7, [meta["period_labels"][p] for p in ACTUAL_MONTHS] + ["Year to date"])
    # The year-to-date column is the sum of the eight beside it: washed as the roll-up.
    column_group(ws, 7, FIRST_DATA_COL + len(ACTUAL_MONTHS), FIRST_DATA_COL + len(ACTUAL_MONTHS))
    r = 8
    ytd_col = col(FIRST_DATA_COL + len(ACTUAL_MONTHS))
    for field, name, sub in meta["cf_rows"]:
        put(ws, f"B{r}", name, "ns_label_sub" if sub else "ns_label")
        for i, p in enumerate(ACTUAL_MONTHS):
            c = col(FIRST_DATA_COL + i)
            ws[f"{c}{r}"] = f"=SUMIFS(cf_{field},cf_period,{p})"
            ws[f"{c}{r}"].style = "ns_m1_sub" if sub else "ns_m1"
        if field == "opening_m":
            ws[f"{ytd_col}{r}"] = f"=SUMIFS(cf_{field},cf_period,{ACTUAL_MONTHS[0]})"
        elif field == "closing_m":
            ws[f"{ytd_col}{r}"] = f"=SUMIFS(cf_{field},cf_period,{REPORT_PERIOD})"
        else:
            ws[f"{ytd_col}{r}"] = (f"=SUM({col(FIRST_DATA_COL)}{r}:"
                                   f"{col(FIRST_DATA_COL + len(ACTUAL_MONTHS) - 1)}{r})")
        ws[f"{ytd_col}{r}"].style = "ns_m1_sub" if sub else "ns_m1"
        r += 1

    put(ws, f"B{r}", "Check — opening plus movements less closing", "ns_label")
    for i, p in enumerate(ACTUAL_MONTHS):
        c = col(FIRST_DATA_COL + i)
        ws[f"{c}{r}"] = meta["cf_check_formula"].format(c=c)
        ws[f"{c}{r}"].style = "ns_m2"
    ws[f"{ytd_col}{r}"] = meta["cf_check_formula"].format(c=ytd_col)
    ws[f"{ytd_col}{r}"].style = "ns_m2"
    r += 2

    section(ws, r, "Liquidity", last_col=11, right_text="At each month end")
    r += 1
    headers(ws, r, [meta["period_labels"][p] for p in ACTUAL_MONTHS] + [""], label_text="")
    r += 1
    for field, name, sub in (("closing_m", "Cash and cash equivalents", False),
                             ("rcf_drawn_m", "Revolving facility drawn", False),
                             ("rcf_avail_m", "Revolving facility available", False),
                             ("liquidity_m", "Total liquidity", True)):
        put(ws, f"B{r}", name, "ns_label_sub" if sub else "ns_label")
        for i, p in enumerate(ACTUAL_MONTHS):
            c = col(FIRST_DATA_COL + i)
            ws[f"{c}{r}"] = f"=SUMIFS(cf_{field},cf_period,{p})"
            ws[f"{c}{r}"].style = "ns_m1_sub" if sub else "ns_m1"
        r += 1
    # The policy floor is a term, not a result -- the same copper the covenant limit carries
    # on sheet 09, for the same reason. Whether cash is above it is read from the row above.
    put(ws, f"B{r}", "Minimum cash policy", "ns_label_copper_i")
    for i in range(len(ACTUAL_MONTHS)):
        c = col(FIRST_DATA_COL + i)
        ws[f"{c}{r}"] = meta["min_cash_policy"]
        ws[f"{c}{r}"].style = "ns_m1_copper"
    r += 2

    section(ws, r, "Trend", last_col=11)
    slots, chart_w = chart_slots(ws, 11)
    line_chart(ws, f"{slots[0]}{r + 2}", "Cash and total liquidity — FY2026",
               Reference(wb["_chart"], min_col=2, min_row=2, max_row=13),
               [(Reference(wb["_chart"], min_col=16, min_row=2, max_row=13), "Cash",
                 S.ACTUAL, None),
                (Reference(wb["_chart"], min_col=17, min_row=2, max_row=13),
                 "Total liquidity", S.FORECAST, None)], width=chart_w, height=7.4,
               page_break=True)
    # Investing is the bar a portfolio owner reads first -- it is the capital going into the
    # business -- so it carries the copper. Operating and financing stay navy. Copper here
    # is a category, not a judgement: an investing outflow is neither good nor bad.
    bar_chart(ws, f"{slots[1]}{r + 2}", "Year-to-date cash flow by category",
              Reference(wb["_chart"], min_col=19, min_row=2, max_row=5),
              [(Reference(wb["_chart"], min_col=20, min_row=2, max_row=5), "USD m",
                S.ACTUAL)], width=chart_w, height=7.4, point_colours={1: S.COPPER})
    note(ws, r + 18, "The statement is derived from balance sheet movements, so it ties by "
                     "construction — the check row above is nil in every month. The effect of "
                     "exchange rates on cash is the retranslation of foreign-currency cash "
                     "and is not the cumulative translation adjustment, which is an equity "
                     "reserve shown on sheet 12.", last_col=11)
    ws.print_area = f"A1:L{r + 20}"
    return ws


# ------------------------------------------------------------------ 07 Working capital
def working_capital(wb, meta):
    ws = wb.create_sheet("07 Working Capital")
    S.sheet_setup(ws, freeze="C8", zoom=100)
    std_widths(ws, label_width=38, data_width=10.5, n_data=9)

    title(ws, "Working capital",
          f"FY{REPORT_FY} to {meta['report_label']} · USD millions and days", width_cols=11)

    section(ws, 5, "Net working capital", last_col=11, right_text="Month end positions")
    headers(ws, 7, [meta["period_labels"][p] for p in ACTUAL_MONTHS])
    r = 8
    for field, name, sub, fmt in (("ar_m", "Accounts receivable", False, "ns_m1"),
                                  ("inv_m", "Inventory", False, "ns_m1"),
                                  ("ap_m", "Accounts payable", False, "ns_m1"),
                                  ("other_m", "Other working capital", False, "ns_m1"),
                                  ("nwc_m", "Net working capital", True, "ns_m1_sub")):
        put(ws, f"B{r}", name, "ns_label_sub" if sub else "ns_label")
        for i, p in enumerate(ACTUAL_MONTHS):
            c = col(FIRST_DATA_COL + i)
            ws[f"{c}{r}"] = f"=SUMIFS(wc_{field},wc_period,{p})"
            ws[f"{c}{r}"].style = fmt
        r += 1
    r += 1

    section(ws, r, "Operational measures", last_col=11,
            right_text="Trailing three months, annualised")
    r += 1
    headers(ws, r, [meta["period_labels"][p] for p in ACTUAL_MONTHS], label_text="")
    r += 1
    for field, name, sub in (("dso_days", "Days sales outstanding", False),
                             ("dio_days", "Days inventory outstanding", False),
                             ("dpo_days", "Days payable outstanding", False),
                             ("ccc_days", "Cash conversion cycle", True)):
        if field == "ccc_days":
            # The cycle is the one derived measure on this page and the one drawn in copper
            # below; the key square ties the row to its chart.
            marker(ws, f"A{r}", "copper")
        put(ws, f"B{r}", name, "ns_label_sub" if sub else "ns_label")
        for i, p in enumerate(ACTUAL_MONTHS):
            c = col(FIRST_DATA_COL + i)
            ws[f"{c}{r}"] = f"=SUMIFS(wc_{field},wc_period,{p})"
            ws[f"{c}{r}"].style = "ns_days"
        r += 1
    r += 1

    note(ws, r, "Days are computed on the trailing three months of revenue and cost of sales, "
                "annualised, because a single month of a project business is not a rate of "
                "activity. They are derived from the consolidated statements only — no "
                "invoice-level ageing exists in the model, so no ageing profile is shown.",
         last_col=11)
    r += 2

    section(ws, r, "Trend", last_col=11)
    slots, chart_w = chart_slots(ws, 11)
    line_chart(ws, f"{slots[0]}{r + 2}", "Net working capital — FY2026",
               Reference(wb["_chart"], min_col=2, min_row=2, max_row=13),
               [(Reference(wb["_chart"], min_col=22, min_row=2, max_row=13),
                 "Net working capital", S.ACTUAL, None)], width=chart_w, height=7.4,
               page_break=True)
    line_chart(ws, f"{slots[1]}{r + 2}", "Cash conversion cycle — days",
               Reference(wb["_chart"], min_col=2, min_row=2, max_row=13),
               [(Reference(wb["_chart"], min_col=23, min_row=2, max_row=13),
                 "Cash conversion cycle", S.COPPER, None)], width=chart_w, height=7.4,
               number_format=S.DAYS)
    ws.print_area = f"A1:L{r + 18}"
    return ws


# ------------------------------------------------------------------ 08 EBITDA bridge
def ebitda_bridge(wb, meta):
    ws = wb.create_sheet("08 EBITDA Bridge")
    S.sheet_setup(ws, freeze="C8", zoom=100)
    std_widths(ws, label_width=48, data_width=13, n_data=5)
    ws.column_dimensions["G"].width = 20   # the covenant-treatment column carries words

    title(ws, "EBITDA bridges",
          "Statutory to management adjusted, and separately to covenant · USD millions",
          width_cols=8)

    section(ws, 5, "Statutory EBITDA to Management Adjusted EBITDA", last_col=8,
            right_text="Approved add-back policy — ADR-0013")
    headers(ws, 7, [f"FY{y}" for y in meta["fy_list"]], label_text="")
    # The reporting year is the column the rest of the pack is about; the three prior years
    # are context. Washed on each of the three bridges so the eye lands on the same column
    # every time.
    current_fy = FIRST_DATA_COL + len(meta["fy_list"]) - 1
    column_group(ws, 7, current_fy, current_fy)
    r = 8
    # On a bridge the two ends are the ledger and the definition; the rows between them are
    # the adjustments -- the same copper the bridge chart gives its middle bar, and the same
    # copper sheet 13 gives the adjustment layers.
    for field, name, sub in (("statutory_m", "Statutory EBITDA", True),
                             ("addbacks_m", "Approved add-backs", False),
                             ("layer4_m", "Management adjustments (layer 4)", False),
                             ("adjusted_m", "Management Adjusted EBITDA", True)):
        if not sub:
            marker(ws, f"A{r}", "copper")
        put(ws, f"B{r}", name, "ns_label_sub" if sub else "ns_label")
        for i, y in enumerate(meta["fy_list"]):
            c = col(FIRST_DATA_COL + i)
            ws[f"{c}{r}"] = f"=SUMIFS(eb_{field},eb_year,{y})"
            ws[f"{c}{r}"].style = "ns_m1_sub" if sub else "ns_m1"
        r += 1
    r += 1

    section(ws, r, "Add-backs by approved category", last_col=8,
            right_text="Inside operating expenses at layer 1")
    r += 1
    headers(ws, r, [f"FY{y}" for y in meta["fy_list"]] + ["Covenant"],
            label_text="Account")
    column_group(ws, r, current_fy, current_fy)
    r += 1
    first = r
    for account, name in meta["addback_accounts"]:
        put(ws, f"B{r}", f"{account}  {name}", "ns_label")
        for i, y in enumerate(meta["fy_list"]):
            c = col(FIRST_DATA_COL + i)
            ws[f"{c}{r}"] = f'=SUMIFS(ab_amount,ab_account,"{account}",ab_year,{y})'
            ws[f"{c}{r}"].style = "ns_m1"
        put(ws, f"{col(FIRST_DATA_COL + len(meta['fy_list']))}{r}",
            "Permitted, capped" if account == "630400" else "Permitted", "ns_text_mut")
        r += 1
    put(ws, f"B{r}", "Total add-backs", "ns_label_sub")
    for i in range(len(meta["fy_list"])):
        c = col(FIRST_DATA_COL + i)
        ws[f"{c}{r}"] = f"=SUM({c}{first}:{c}{r-1})"
        ws[f"{c}{r}"].style = "ns_m1_sub"
    r += 2

    section(ws, r, "Statutory EBITDA to Covenant EBITDA", last_col=8,
            right_text="Credit agreement CA-021 to CA-030")
    r += 1
    headers(ws, r, [f"FY{y}" for y in meta["fy_list"]], label_text="")
    column_group(ws, r, current_fy, current_fy)
    r += 1
    for field, name, sub in (("statutory_m", "Statutory EBITDA", True),
                             ("addbacks_m", "Permitted add-backs", False),
                             ("cap_effect_m", "Sponsor fee cap effect (CA-027)", False),
                             ("fx_addback_m", "Unrealised foreign exchange (CA-030)", False),
                             ("covenant_m", "Covenant EBITDA", True)):
        if not sub:
            marker(ws, f"A{r}", "copper")
        put(ws, f"B{r}", name, "ns_label_sub" if sub else "ns_label")
        for i, y in enumerate(meta["fy_list"]):
            c = col(FIRST_DATA_COL + i)
            ws[f"{c}{r}"] = f"=SUMIFS(eb_{field},eb_year,{y})"
            ws[f"{c}{r}"].style = "ns_m1_sub" if sub else "ns_m1"
        r += 1
    r += 1

    section(ws, r, "Why the two measures are not the same", last_col=8)
    r += 1
    headers(ws, r, ["Management Adjusted", "", "Covenant"], label_text="Treatment")
    ws.cell(row=r, column=3).alignment = ws.cell(row=r, column=3).alignment.copy(
        horizontal="left", indent=1)
    ws.cell(row=r, column=5).alignment = ws.cell(row=r, column=5).alignment.copy(
        horizontal="left", indent=1)
    r += 1
    for item, mgmt, cov in meta["ebitda_policy"]:
        put(ws, f"B{r}", item, "ns_label")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        put(ws, f"C{r}", mgmt, "ns_text_mut")
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
        put(ws, f"E{r}", cov, "ns_text_mut")
        r += 1
    note(ws, r + 1, "The two measures are equal in every year presented, and that is an "
                    "outcome rather than a definition: the sponsor monitoring fee runs below "
                    "the CA-027 cap so the cap does not bite, and the CA-030 unrealised "
                    "foreign exchange add-back has no population in this window. The cap is "
                    "still computed from the agreement's own term, because it would bite at a "
                    "higher fee and the calculation must not need editing when it does.",
         last_col=8)
    r += 3

    # The bridge's middle components -- the approved add-backs and the layer-4 management
    # effect -- are what turns statutory EBITDA into the management measure. They are given
    # copper so a reader can see at a glance which bars are the adjustment and which are the
    # two EBITDA definitions either side of it. Copper here means "management adjustment",
    # never favourable or unfavourable; those stay green and red and mean what they always did.
    # Width derived from the sheet, not fixed. At a hardcoded 16.5 cm this chart ran past the
    # print area's right edge and Excel simply did not render it -- the page printed with an
    # empty band where the bridge should have been, and nothing reported an error.
    _, bridge_w = chart_slots(ws, 8, count=1)
    bar_chart(ws, f"B{r}", "EBITDA bridge — FY2026 statutory to adjusted",
              Reference(wb["_chart"], min_col=25, min_row=2, max_row=5),
              [(Reference(wb["_chart"], min_col=26, min_row=2, max_row=5), "USD m",
                S.ACTUAL)], width=bridge_w, height=7.4,
              point_colours={1: S.COPPER, 2: S.COPPER})
    ws.print_area = f"A1:H{r + 16}"
    return ws


# ------------------------------------------------------------------ 09 Debt and covenants
def debt_covenants(wb, meta):
    ws = wb.create_sheet("09 Debt & Covenants")
    S.sheet_setup(ws, freeze="C8", zoom=100)
    S.widths(ws, {"A": 2.2, "B": 36, "C": 13, "D": 10, "E": 11, "F": 11, "G": 20,
                  "H": 13, "I": 9, "J": 12, "K": 13})

    title(ws, "Debt and covenant compliance",
          f"At {meta['report_label']} · USD millions", width_cols=11)

    section(ws, 5, "Debt by instrument", last_col=11, right_text="At the reporting date")
    headers(ws, 7, ["Type", "Currency", "Drawn", "Undrawn", "Rate", "Fixed / floating",
                    "Hedged", "Maturity", "In covenant debt"], label_text="Instrument")
    # The last column is the agreement's definition applied to each instrument -- a term,
    # like the limit below, and washed as one.
    column_group(ws, 7, FIRST_DATA_COL + 8, FIRST_DATA_COL + 8)
    r = 8
    first = r
    for kind, key, iname, itype, ccy, rate, rtype, hedged, maturity, in_cov in \
            meta["instruments"]:
        # One row may be a single instrument or a whole class of them; the criterion column
        # changes, the arithmetic does not, and the total still sums the rows on the page.
        criterion = "dbt_id" if kind == "id" else "dbt_type"
        put(ws, f"B{r}", iname, "ns_label")
        put(ws, f"C{r}", itype.replace("_", " ").title(), "ns_text_mut")
        put(ws, f"D{r}", ccy, "ns_text_c")
        ws[f"E{r}"] = (f'=SUMIFS(dbt_closing,{criterion},"{key}",'
                       f'dbt_period,{REPORT_PERIOD})')
        ws[f"E{r}"].style = "ns_m1"
        ws[f"F{r}"] = (f'=SUMIFS(dbt_undrawn,{criterion},"{key}",'
                       f'dbt_period,{REPORT_PERIOD})')
        ws[f"F{r}"].style = "ns_m1"
        put(ws, f"G{r}", rate, "ns_text_c")
        put(ws, f"H{r}", rtype.title(), "ns_text_c")
        put(ws, f"I{r}", "Yes" if hedged else "No", "ns_text_c")
        put(ws, f"J{r}", maturity, "ns_text_c")
        put(ws, f"K{r}", "Yes" if in_cov else "No", "ns_text_c")
        r += 1
    put(ws, f"B{r}", "Total debt", "ns_label_sub")
    for c in "EF":
        ws[f"{c}{r}"] = f"=SUM({c}{first}:{c}{r-1})"
        ws[f"{c}{r}"].style = "ns_m1_sub"
    r += 2

    section(ws, r, "Net leverage covenant", last_col=11,
            right_text="Tested at each fiscal year end")
    r += 1
    headers(ws, r, [meta["period_labels"][p] for p in meta["cov_periods"]], label_text="")
    # Three test dates and one that is not: the current month is shown for management
    # information, and the wash separates it from the columns the agreement is tested on.
    column_group(ws, r, FIRST_DATA_COL + len(meta["cov_periods"]) - 1,
                 FIRST_DATA_COL + len(meta["cov_periods"]) - 1)
    r += 1
    for field, name, sub, fmt in (
            ("covenant_debt_m", "Covenant debt", False, "ns_m1"),
            ("cash_m", "Less cash", False, "ns_m1"),
            ("net_debt_m", "Net debt", True, "ns_m1_sub"),
            ("covenant_ebitda_m", "Covenant EBITDA, last twelve months", False, "ns_m1"),
            ("net_leverage", "Net leverage", True, "ns_ratio_sub"),
            ("max_net_leverage", "Covenant limit", False, "ns_ratio_copper"),
            ("headroom_turns", "Headroom", True, "ns_ratio_sub"),
            ("headroom_m", "Headroom in EBITDA terms", False, "ns_m1"),
            ("economic_leverage", "Economic leverage, with operating leases",
             False, "ns_ratio")):
        put(ws, f"B{r}", name, "ns_label_sub" if sub else "ns_label")
        for i, p in enumerate(meta["cov_periods"]):
            c = col(FIRST_DATA_COL + i)
            sign = "-" if field == "cash_m" else ""
            ws[f"{c}{r}"] = f"={sign}SUMIFS(cov_{field},cov_period,{p})"
            ws[f"{c}{r}"].style = fmt
        if field == "headroom_turns":
            head_row = r
        r += 1

    put(ws, f"B{r}", "Status", "ns_label_sub")
    for i, p in enumerate(meta["cov_periods"]):
        c = col(FIRST_DATA_COL + i)
        if p == meta["cov_indicative"] and p % 100 != 12:
            # not a test date under the agreement, so it does not get a compliance verdict
            ws[f"{c}{r}"] = "Indicative"
            ws[f"{c}{r}"].style = "ns_status_info"
        else:
            ws[f"{c}{r}"] = f'=IF({c}{head_row}>=0,"Compliant","BREACH")'
            ws[f"{c}{r}"].style = "ns_status_ok"
    ws.conditional_formatting.add(
        f"{col(FIRST_DATA_COL)}{r}:{col(FIRST_DATA_COL + len(meta['cov_periods']) - 1)}{r}",
        __import__("openpyxl.formatting.rule", fromlist=["CellIsRule"]).CellIsRule(
            operator="equal", formula=['"BREACH"'],
            font=Font(name=S.FONT, bold=True, color=S.UNFAVOURABLE)))
    r += 2

    note(ws, r, "The agreement sets a maximum leverage per fiscal year — 6.00x FY2023, 5.50x "
                "FY2024, 5.00x FY2025 and 4.50x from FY2026 — so the test dates are the year "
                "ends, and each limit is read from its own agreement term rather than written "
                "into the report. The current month is shown for management information and "
                "is not a test date. Leverage is measured on twelve months of covenant "
                "EBITDA against net debt at the date. Economic leverage adds operating lease "
                "liabilities, which the agreement excludes (CA-018), and is shown because a "
                "reader should see both.", last_col=11)
    r += 2

    section(ws, r, "Trend", last_col=11)
    slots, chart_w = chart_slots(ws, 11)
    line_chart(ws, f"{slots[0]}{r + 2}", "Net leverage against the covenant limit",
               Reference(wb["_chart"], min_col=2, min_row=2, max_row=13),
               [(Reference(wb["_chart"], min_col=28, min_row=2, max_row=13), "Net leverage",
                 S.ACTUAL, None),
                # The limit is a term of the agreement, not a warning. In red it read as an
                # alarm on a chart where nothing was wrong; in copper it reads as the
                # reference line it is. Breach, where it happens, is still red -- in the
                # Status row, where a status belongs.
                (Reference(wb["_chart"], min_col=29, min_row=2, max_row=13), "Covenant limit",
                 S.COPPER, "dash"),
                (Reference(wb["_chart"], min_col=30, min_row=2, max_row=13),
                 "Economic leverage", S.BUDGET, None)],
               width=chart_w, height=7.4, number_format=S.RATIO, y_min=3.0, y_max=5.0,
               page_break=True)
    bar_chart(ws, f"{slots[1]}{r + 2}", "Drawn debt by instrument type",
              Reference(wb["_chart"], min_col=32, min_row=2, max_row=5),
              [(Reference(wb["_chart"], min_col=33, min_row=2, max_row=5), "USD m",
                S.ACTUAL)], width=chart_w, height=7.4)
    ws.print_area = f"A1:K{r + 18}"
    return ws
